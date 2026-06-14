from langchain.tools import tool
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.dml.color import RGBColor
from docx.oxml.shared import OxmlElement
import os
import platform

# ======================
# 【全局配置】
# ======================
TITLE_FONT = "黑体"
BODY_FONT = "宋体"
TITLE_SIZE = 22
H1_SIZE = 18
H2_SIZE = 16
BODY_SIZE = 14
INDENT = 2
LINE_SPACING = 1.5
SPACE_BEFORE = 4
SPACE_AFTER = 6
FONT_COLOR = (0, 0, 0)
CHAR_SPACING = 0

MARGIN_LEFT = 2.5
MARGIN_RIGHT = 2.5
MARGIN_TOP = 2.5
MARGIN_BOTTOM = 2.5

ESSAY_MARGIN = 2.0
ESSAY_INDENT = 2
ESSAY_LINE = 1.2


def get_desktop():
    return os.path.join(os.path.expanduser("~"), "Desktop")


def set_font(run, font_name, size, bold=False, italic=False, color=(0, 0, 0)):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color)


def set_char_spacing(run, value):
    rPr = run._element.get_or_add_rPr()
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:val'), str(int(value * 20)))
    rPr.append(spacing)


# ======================
# 【修复版 TOC 目录】
# ======================
def add_toc(doc):
    p_title = doc.add_paragraph()
    run = p_title.add_run("目录")
    set_font(run, TITLE_FONT, 16, bold=True)
    p_title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    doc.add_paragraph()

    toc_p = doc.add_paragraph()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')

    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = r'TOC \o "1-2" \h \z \u'

    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')

    placeholder_run = toc_p.add_run("按 Ctrl+A → F9 刷新目录")
    placeholder_run.font.color.rgb = RGBColor(128, 128, 128)
    placeholder_run.italic = True

    toc_p._element.append(fldChar1)
    toc_p._element.append(instrText)
    toc_p._element.append(fldChar2)


# ======================
# 【自动更新目录】
# ======================
def update_toc_automatically(file_path):
    """自动更新Word文档目录"""
    # Windows 自动更新
    if platform.system() == "Windows":
        try:
            import win32com.client  # noqa
            word = win32com.client.Dispatch("Word.Application")
            word.Visible = False
            doc = word.Documents.Open(file_path)
            if doc.TablesOfContents.Count > 0:
                doc.TablesOfContents(1).Update()
            doc.Save()
            doc.Close()
            word.Quit()
            return " 目录已自动更新（Windows Word）"
        except:
            pass

    # 跨平台备选
    try:
        import aspose.words as aw  # noqa
        doc = aw.Document(file_path)
        doc.update_fields()
        doc.save(file_path)
        return " 目录已自动更新（Aspose.Words）"
    except:
        pass

    return "⚠ 请手动刷新：打开文档 → Ctrl+A → F9"


# ======================
# 【1. 专业论文工具】
# ======================
@tool
def create_professional_word(filename: str, title: str, content: str):
    """
    生成带自动目录的专业Word论文
    参数:
        filename: 文档文件名（无需加.docx）
        title: 论文主标题
        content: 论文正文内容，支持一、二级标题自动识别
    返回: 生成结果提示
    """
    try:
        doc = Document()
        sec = doc.sections[0]
        sec.left_margin = Cm(MARGIN_LEFT)
        sec.right_margin = Cm(MARGIN_RIGHT)
        sec.top_margin = Cm(MARGIN_TOP)
        sec.bottom_margin = Cm(MARGIN_BOTTOM)

        # 主标题
        title_p = doc.add_paragraph()
        title_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        r = title_p.add_run(title)
        set_font(r, TITLE_FONT, TITLE_SIZE, bold=True)
        set_char_spacing(r, CHAR_SPACING)
        doc.add_paragraph()

        add_toc(doc)
        doc.add_page_break()

        # 正文解析
        lines = content.strip().split("\n")
        for line in lines:
            line = line.strip()
            if not line:
                continue
            if line.startswith(("一、", "1、", "1.")):
                p = doc.add_paragraph(style="Heading 1")
                run = p.add_run(line)
                set_font(run, TITLE_FONT, H1_SIZE, bold=True)
                p.paragraph_format.space_before = Pt(SPACE_BEFORE)
                p.paragraph_format.space_after = Pt(SPACE_AFTER)
                continue
            if line.startswith(("（1）", "①", "1.1", "2.1")):
                p = doc.add_paragraph(style="Heading 2")
                run = p.add_run(line)
                set_font(run, TITLE_FONT, H2_SIZE, bold=True)
                continue
            p = doc.add_paragraph(line)
            p.paragraph_format.first_line_indent = Cm(INDENT)
            p.paragraph_format.line_spacing = LINE_SPACING
            p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
            for run in p.runs:
                set_font(run, BODY_FONT, BODY_SIZE, color=FONT_COLOR)
                set_char_spacing(run, CHAR_SPACING)

        desktop = get_desktop()
        if not filename.endswith(".docx"):
            filename += ".docx"
        save_path = os.path.join(desktop, filename)
        doc.save(save_path)

        update_msg = update_toc_automatically(save_path)
        return f" 论文生成成功：{save_path}\n{update_msg}"

    except Exception as e:
        return f" 生成失败：{str(e)}"


# ======================
# 【2. 考场作文工具】
# ======================
@tool
def create_essay_word(filename: str, title: str, content: str):
    """
    生成简洁格式的考场作文Word文档
    参数:
        filename: 文档文件名（无需加.docx）
        title: 作文标题
        content: 作文正文内容
    返回: 生成结果提示
    """
    try:
        doc = Document()
        sec = doc.sections[0]
        sec.left_margin = Cm(ESSAY_MARGIN)
        sec.right_margin = Cm(ESSAY_MARGIN)
        sec.top_margin = Cm(ESSAY_MARGIN)
        sec.bottom_margin = Cm(ESSAY_MARGIN)

        # 标题
        title_p = doc.add_paragraph()
        title_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        run = title_p.add_run(title)
        set_font(run, TITLE_FONT, TITLE_SIZE, bold=True)
        doc.add_paragraph()

        # 正文
        lines = content.strip().split("\n")
        for line in lines:
            line = line.strip()
            if not line:
                continue
            p = doc.add_paragraph(line)
            p.paragraph_format.first_line_indent = Cm(ESSAY_INDENT)
            p.paragraph_format.line_spacing = ESSAY_LINE
            p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            for run in p.runs:
                set_font(run, BODY_FONT, BODY_SIZE, color=FONT_COLOR)

        desktop = get_desktop()
        if not filename.endswith(".docx"):
            filename += ".docx"
        save_path = os.path.join(desktop, filename)
        doc.save(save_path)
        return f" 作文生成成功：{save_path}"

    except Exception as e:
        return f" 作文生成失败：{str(e)}"